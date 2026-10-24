import io
import os

import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches

# Settings are read from the environment; tests never talk to Cloudflare or production.
os.environ.setdefault("DATABASE_URL", os.environ.get("TEST_DATABASE_URL", "postgresql://unused/unused"))
os.environ.setdefault("CLOUDFLARE_ACCOUNT_ID", "test-account")
os.environ.setdefault("CLOUDFLARE_API_TOKEN", "test-token")
os.environ.setdefault("RAG_SERVICE_SECRET", "test-secret")
# OCR calls a paid/quota'd vision API; tests that need it stub the provider and switch it on themselves.
os.environ.setdefault("OCR_PROVIDER", "off")

LEFT = (
    "LEFT COLUMN: searching a B-tree costs O(log n) disk reads because each node "
    "holds many keys and the tree stays balanced after every insertion."
)
RIGHT = (
    "RIGHT COLUMN: deletion may merge a node with a sibling or borrow a key from it "
    "so that every node keeps at least the minimum number of keys."
)


def _footer(page: pymupdf.Page, n: int) -> None:
    page.insert_text((72, 800), f"CS201 Data Structures - Lecture Notes - Page {n}", fontsize=8)


@pytest.fixture(scope="session")
def lecture_pdf() -> bytes:
    """3 pages: heading + two columns, a second heading, and a page with only a footer."""
    doc = pymupdf.open()
    p1 = doc.new_page(width=595, height=842)
    p1.insert_text((72, 80), "Unit 3: B-Trees", fontsize=22)
    p1.insert_textbox(pymupdf.Rect(72, 110, 523, 160),
                      "A B-tree of order m keeps every node between ceil(m/2) and m children.", fontsize=11)
    p1.insert_textbox(pymupdf.Rect(72, 180, 287, 400), LEFT, fontsize=11)
    p1.insert_textbox(pymupdf.Rect(308, 180, 523, 400), RIGHT, fontsize=11)
    _footer(p1, 1)

    p2 = doc.new_page(width=595, height=842)
    p2.insert_text((72, 80), "Unit 4: Hashing", fontsize=22)
    p2.insert_textbox(pymupdf.Rect(72, 110, 523, 300),
                      "Open addressing resolves collisions by probing. Linear probing suffers from "
                      "primary clustering, which quadratic probing reduces.", fontsize=11)
    _footer(p2, 2)

    p3 = doc.new_page(width=595, height=842)  # e.g. a scanned page: no text layer
    _footer(p3, 3)
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture(scope="session")
def picture_pptx() -> bytes:
    """4 slides: text; title + picture; short text only (hidden); picture only, with speaker notes."""
    png = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 30), 0).tobytes("png")
    prs = Presentation()

    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Kerberos"
    s1.placeholders[1].text_frame.text = "Network authentication protocol developed at MIT"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    s2.shapes.title.text = "Waiting"
    s2.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(2), Inches(4), Inches(3))

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Thank you"
    s3._element.set("show", "0")  # hidden: LibreOffice must still export it, or slide numbers shift

    s4 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    s4.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(1), Inches(6), Inches(4))
    s4.notes_slide.notes_text_frame.text = "Walk through the LSTM gates on this diagram."

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def lecture_pptx() -> bytes:
    """Slides: two short slides sharing a title, a table slide, and a slide with notes."""
    prs = Presentation()
    layout = prs.slide_layouts[1]  # Title and Content

    s1 = prs.slides.add_slide(layout)
    s1.shapes.title.text = "Kerberos"
    s1.placeholders[1].text_frame.text = "Network authentication protocol"

    s2 = prs.slides.add_slide(layout)
    s2.shapes.title.text = "Kerberos"
    s2.placeholders[1].text_frame.text = "Developed at MIT in the mid 1980s"

    s3 = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    s3.shapes.title.text = "LSTM vs GRU"
    table = s3.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, row in enumerate([["Feature", "LSTM", "GRU"], ["Gates", "3", "2"], ["Training speed", "Slower", "Faster"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    s4 = prs.slides.add_slide(layout)
    s4.shapes.title.text = "Ticket Granting Ticket"
    s4.placeholders[1].text_frame.text = (
        "Instead of sending the password to every server, the client requests a ticket "
        "granting ticket once and uses it to obtain service tickets for each application."
    )
    s4.notes_slide.notes_text_frame.text = "Mention that a stolen TGT only works until it expires."

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
