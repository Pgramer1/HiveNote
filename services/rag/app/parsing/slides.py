import io
import shutil
import subprocess
import tempfile
from pathlib import Path
from zipfile import BadZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

from app.cleaning import normalize_whitespace
from app.parsing.model import Block, ParsedDocument, UnsupportedDocument


def _shape_lines(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [line for child in shape.shapes for line in _shape_lines(child)]
    if getattr(shape, "has_table", False) and shape.has_table:
        return [" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
    return []


def _has_picture(shape) -> bool:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return any(_has_picture(child) for child in shape.shapes)
    return shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image")


# Same threshold as PDFs (app/parsing/pdf.py): below it the slide's content is probably in a picture.
_MIN_SLIDE_CHARS = 20


def _reading_order(shape) -> tuple[int, int]:
    return (shape.top or 0, shape.left or 0)


def parse_pptx(data: bytes) -> ParsedDocument:
    try:
        prs = Presentation(io.BytesIO(data))
    except (BadZipFile, PackageNotFoundError, KeyError, ValueError) as exc:
        raise UnsupportedDocument(f"Not a readable .pptx file: {exc}") from exc

    blocks: list[Block] = []
    empty_pages: list[int] = []
    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.text.strip() else None

        lines = [title] if title else []
        for shape in sorted(slide.shapes, key=_reading_order):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            lines.extend(_shape_lines(shape))
        text = normalize_whitespace("\n".join(lines))

        letters = sum(ch.isalnum() for ch in text)
        # Little text plus a picture: the content is in the picture, so OCR reads the slide. The short
        # text (usually the title) is kept until OCR replaces it. Short text-only slides ("Thank you")
        # aren't worth an OCR call.
        if letters < _MIN_SLIDE_CHARS and any(_has_picture(shape) for shape in slide.shapes):
            empty_pages.append(number)
        if letters >= 3:
            blocks.append(Block(text=text, page=number, section=title, kind="slide"))

        if slide.has_notes_slide:
            notes = normalize_whitespace(slide.notes_slide.notes_text_frame.text or "")
            if notes:
                blocks.append(Block(text=f"Speaker notes: {notes}", page=number, section=title, kind="notes"))

    return ParsedDocument(blocks=blocks, page_count=len(prs.slides), empty_pages=empty_pages, slides=True)


def _soffice_convert(data: bytes, convert_to: str, out_ext: str, timeout: int) -> bytes:
    """Runs LibreOffice (installed in the Docker image and the GitHub workflow) on a deck."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise UnsupportedDocument("Converting decks needs LibreOffice (soffice), which isn't installed")

    with tempfile.TemporaryDirectory() as tmp:
        # LibreOffice picks the import filter from the extension: .pptx is a zip, legacy .ppt is OLE.
        src = Path(tmp) / ("deck.pptx" if data[:4] == b"PK\x03\x04" else "deck.ppt")
        src.write_bytes(data)
        out = Path(tmp) / "out"
        profile = (Path(tmp) / "profile").as_uri()  # isolated profile: safe to run concurrently
        result = subprocess.run(
            [soffice, f"-env:UserInstallation={profile}", "--headless", "--norestore",
             "--convert-to", convert_to, "--outdir", str(out), str(src)],
            capture_output=True, timeout=timeout,
        )
        converted = out / f"deck.{out_ext}"
        if result.returncode != 0 or not converted.exists():
            detail = (result.stderr or result.stdout).decode(errors="replace")[-300:]
            raise UnsupportedDocument(f"LibreOffice couldn't convert the deck to {out_ext}: {detail}")
        return converted.read_bytes()


def convert_ppt_to_pptx(data: bytes, timeout: int = 180) -> bytes:
    """Converts a legacy binary .ppt to .pptx."""
    return _soffice_convert(data, "pptx", "pptx", timeout)


# Hidden slides are exported too, so PDF page N is slide N as python-pptx counts them.
_PDF_EXPORT = 'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"}}'


def convert_deck_to_pdf(data: bytes, timeout: int = 180) -> bytes:
    """Renders a .pptx or .ppt deck to PDF, one page per slide, so OCR can read slide images."""
    return _soffice_convert(data, _PDF_EXPORT, "pdf", timeout)
